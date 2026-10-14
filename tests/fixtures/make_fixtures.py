"""Generate synthetic invoice fixtures deterministically for testing.

Run: uv run python tests/fixtures/make_fixtures.py
Output: tests/fixtures/generated/
"""

from pathlib import Path

GENERATED_DIR = Path(__file__).parent / "generated"


def make_utf8_txt() -> None:
    content = """INVOICE

Vendor: Acme Corp
GSTIN: 29ABCDE1234F1Z5
Address: 123 Main St, Bangalore

Invoice Number: INV-2026-001
Invoice Date: 2026-01-15
Due Date: 2026-02-15

Bill To:
Buyer Inc
GSTIN: 07XYZAB5678C1D2

| Description | HSN  | Qty | Unit | Unit Price | Amount   |
|-------------|------|-----|------|------------|----------|
| Widget A    | 8471 | 10  | pcs  | 100.00     | 1000.00  |
| Widget B    | 8472 | 5   | pcs  | 36.00      | 180.00   |

Subtotal: 1180.00
CGST @9%: 106.20
SGST @9%: 106.20
Total: 1392.40
Currency: INR
"""
    (GENERATED_DIR / "invoice_utf8.txt").write_text(content, encoding="utf-8")


def make_latin1_txt() -> None:
    content = "INVOICE\n\nVendor: Café Résumé\nTotal: ₹1180.00\nInvoice Number: INV-LAT-001\n"
    (GENERATED_DIR / "invoice_latin1.txt").write_bytes(content.encode("windows-1252", errors="replace"))


def make_docx_with_table() -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("INVOICE", level=1)
    doc.add_paragraph("Vendor: Acme Corp")
    doc.add_paragraph("Invoice Number: INV-DOCX-001")

    table = doc.add_table(rows=3, cols=4)
    table.style = "Table Grid"
    headers = ["Description", "Qty", "Unit Price", "Amount"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    rows_data = [
        ["Widget A", "10", "100.00", "1000.00"],
        ["Widget B", "5", "36.00", "180.00"],
    ]
    for r, row_data in enumerate(rows_data, start=1):
        for c, val in enumerate(row_data):
            table.rows[r].cells[c].text = val

    doc.add_paragraph("Total: 1180.00")
    doc.save(str(GENERATED_DIR / "invoice_table.docx"))


def make_docx_with_header() -> None:
    from docx import Document

    doc = Document()
    section = doc.sections[0]
    header = section.header
    header.paragraphs[0].text = "Acme Corp - Official Invoice"
    footer = section.footer
    footer.paragraphs[0].text = "Page 1"

    doc.add_paragraph("Invoice Number: INV-HDR-001")
    doc.add_paragraph("Total: 500.00")
    doc.save(str(GENERATED_DIR / "invoice_header.docx"))


def make_digital_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = str(GENERATED_DIR / "invoice_digital.pdf")
    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("Helvetica", 12)

    c.drawString(72, 750, "INVOICE")
    c.drawString(72, 730, "Vendor: Acme Corp")
    c.drawString(72, 710, "Invoice Number: INV-PDF-001")
    c.drawString(72, 690, "Date: 2026-01-15")
    c.drawString(72, 660, "Widget A    10 x 100.00 = 1000.00")
    c.drawString(72, 640, "Widget B     5 x  36.00 =  180.00")
    c.drawString(72, 610, "Total: 1180.00")
    c.showPage()

    c.setFont("Helvetica", 12)
    c.drawString(72, 750, "Page 2 - Terms and Conditions")
    c.drawString(72, 730, "Payment due within 30 days.")
    c.showPage()
    c.save()


def make_scanned_pdf() -> None:
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    img = Image.new("RGB", (800, 600), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 20)
    except OSError:
        font = ImageFont.load_default()
    draw.text((50, 50), "INVOICE - SCANNED", fill=(0, 0, 0), font=font)
    draw.text((50, 100), "Vendor: Scan Corp", fill=(0, 0, 0), font=font)
    draw.text((50, 150), "Invoice Number: INV-SCAN-001", fill=(0, 0, 0), font=font)
    draw.text((50, 200), "Total: 2500.00", fill=(0, 0, 0), font=font)

    img_path = str(GENERATED_DIR / "_tmp_scanned.png")
    img.save(img_path, format="PNG")

    path = str(GENERATED_DIR / "invoice_scanned.pdf")
    c = canvas.Canvas(path, pagesize=A4)
    c.drawImage(img_path, 0, 0, width=A4[0], height=A4[1], preserveAspectRatio=True)
    c.showPage()
    c.save()
    Path(img_path).unlink()


def make_mixed_pdf() -> None:
    """Page 1 has text, page 2 is image-only."""
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = str(GENERATED_DIR / "invoice_mixed.pdf")
    c = canvas.Canvas(path, pagesize=A4)

    c.setFont("Helvetica", 12)
    c.drawString(72, 750, "Page 1 - Digital text content")
    c.drawString(72, 730, "Vendor: Mixed Corp")
    c.showPage()

    img = Image.new("RGB", (800, 600), color=(245, 245, 245))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 20)
    except OSError:
        font = ImageFont.load_default()
    draw.text((50, 50), "Page 2 - Scanned content", fill=(0, 0, 0), font=font)
    img_path = str(GENERATED_DIR / "_tmp_mixed.png")
    img.save(img_path, format="PNG")

    c.drawImage(img_path, 0, 0, width=A4[0], height=A4[1], preserveAspectRatio=True)
    c.showPage()
    c.save()
    Path(img_path).unlink()


def make_image_fixtures() -> None:
    from PIL import Image, ImageDraw, ImageFont

    for fmt, ext in [("PNG", "png"), ("JPEG", "jpg")]:
        img = Image.new("RGB", (800, 600), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        try:
            font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 20)
        except OSError:
            font = ImageFont.load_default()
        draw.text((50, 50), f"INVOICE - {fmt}", fill=(0, 0, 0), font=font)
        draw.text((50, 100), "Vendor: Image Corp", fill=(0, 0, 0), font=font)
        draw.text((50, 150), "Total: 999.00", fill=(0, 0, 0), font=font)
        img.save(str(GENERATED_DIR / f"invoice.{ext}"), format=fmt)


def make_pdf_renamed_as_txt() -> None:
    """A PDF file saved with .txt extension."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = str(GENERATED_DIR / "actually_a_pdf.txt")
    c = canvas.Canvas(path, pagesize=A4)
    c.setFont("Helvetica", 12)
    c.drawString(72, 750, "This is a PDF with wrong extension")
    c.showPage()
    c.save()


def make_empty_file() -> None:
    (GENERATED_DIR / "empty_file.txt").write_bytes(b"")


def make_zip_file() -> None:
    import zipfile
    path = GENERATED_DIR / "archive.zip"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("dummy.txt", "dummy content")


def make_xlsx_file() -> None:
    (GENERATED_DIR / "spreadsheet.xlsx").write_bytes(
        b"PK\x03\x04fake xlsx content"
    )


def make_encrypted_pdf() -> None:
    import fitz
    path = str(GENERATED_DIR / "encrypted.pdf")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Secret content")
    perm = fitz.PDF_PERM_ACCESSIBILITY
    encrypt_meth = fitz.PDF_ENCRYPT_AES_256
    doc.save(
        path,
        encryption=encrypt_meth,
        owner_pw="owner123",
        user_pw="user123",
        permissions=perm,
    )
    doc.close()


def make_corrupt_pdf() -> None:
    (GENERATED_DIR / "corrupt.pdf").write_bytes(b"%PDF-1.4 corrupt data here \x00\x01\x02")


def make_large_pdf() -> None:
    """PDF with 25 pages (over default limit of 20)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = str(GENERATED_DIR / "large_25pages.pdf")
    c = canvas.Canvas(path, pagesize=A4)
    for i in range(25):
        c.setFont("Helvetica", 12)
        c.drawString(72, 750, f"Page {i + 1}")
        c.showPage()
    c.save()


def make_doc_from_docx() -> None:
    """Convert the docx table fixture to .doc using LibreOffice."""
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        print("  Skipping .doc generation (soffice not found)")
        return

    docx_path = GENERATED_DIR / "invoice_table.docx"
    if not docx_path.exists():
        print("  Skipping .doc generation (invoice_table.docx not found)")
        return

    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "doc", str(docx_path),
             "--outdir", str(GENERATED_DIR)],
            check=True,
            capture_output=True,
            timeout=60,
        )
        print("  Generated invoice_table.doc")
    except (subprocess.CalledProcessError, FileNotFoundError) as e:
        print(f"  Skipping .doc generation (conversion failed: {e})")


def main() -> None:
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    make_utf8_txt()
    make_latin1_txt()
    make_docx_with_table()
    make_docx_with_header()
    make_digital_pdf()
    make_scanned_pdf()
    make_mixed_pdf()
    make_image_fixtures()
    make_pdf_renamed_as_txt()
    make_empty_file()
    make_zip_file()
    make_xlsx_file()
    make_encrypted_pdf()
    make_corrupt_pdf()
    make_large_pdf()
    make_doc_from_docx()
    print(f"Generated fixtures in {GENERATED_DIR}")
    for f in sorted(GENERATED_DIR.iterdir()):
        print(f"  {f.name} ({f.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
