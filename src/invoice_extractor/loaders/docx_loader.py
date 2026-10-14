from __future__ import annotations

from pathlib import Path

from docx import Document

from invoice_extractor.schema import NormalizedDoc


def load_docx(path: Path) -> NormalizedDoc:
    doc = Document(str(path))
    parts: list[str] = []

    for section in doc.sections:
        header = section.header
        if header and header.paragraphs:
            header_text = "\n".join(p.text for p in header.paragraphs if p.text.strip())
            if header_text.strip():
                parts.append(header_text)

    for element in doc.element.body:
        tag = element.tag.split("}")[-1] if "}" in element.tag else element.tag
        if tag == "p":
            from docx.oxml.ns import qn
            texts = [node.text or "" for node in element.iter(qn("w:t"))]
            line = "".join(texts).strip()
            if line:
                parts.append(line)
        elif tag == "tbl":
            from docx.table import Table
            table = Table(element, doc)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append("| " + " | ".join(cells) + " |")

    for section in doc.sections:
        footer = section.footer
        if footer and footer.paragraphs:
            footer_text = "\n".join(p.text for p in footer.paragraphs if p.text.strip())
            if footer_text.strip():
                parts.append(footer_text)

    text = "\n".join(parts)
    return NormalizedDoc(
        mode="text",
        text=text,
        images=None,
        source_format="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        page_count=1,
    )
